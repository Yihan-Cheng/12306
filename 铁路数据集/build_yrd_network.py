# -*- coding: utf-8 -*-
"""
构建长三角铁路列车运行网络
输入：
  railwayCrawler-main/schedulers/*.csv   每个车站的列车停靠记录
  China-rail-way-stations-data-main/src/station.csv  车站->省份/城市 映射
输出：
  output/train.csv  车次运行图（裁剪到长三角范围）
  output/station.csv  车站表
  output/edge.csv  相邻停站边表
  output/长三角铁路运行网络.xlsx  上述三张表 + 说明
"""
import csv, glob, os, sys
from collections import defaultdict

BASE = r"E:\学习\2026秋季学期（大三上）\数据库系统\12306\铁路数据集"
SCHED = os.path.join(BASE, "railwayCrawler-main", "schedulers")
STATION_CSV = os.path.join(BASE, "China-rail-way-stations-data-main", "src", "station.csv")
OUT = os.path.join(BASE, "output")
os.makedirs(OUT, exist_ok=True)

YRD_PROVINCES = {"上海", "江苏", "浙江", "安徽"}
PROV_ORDER = {"上海": 0, "江苏": 1, "浙江": 2, "安徽": 3}

# schedulers 中出现、但 station.csv 缺失的车站 -> 手工补 (省, 市)
MANUAL = {
    "万州北": ("重庆", "重庆"), "中卫南": ("宁夏", "中卫"), "乌兰察布": ("内蒙古", "乌兰察布"),
    "二连": ("内蒙古", "锡林郭勒"), "内江北": ("四川", "内江"), "十堰东": ("湖北", "十堰"),
    "南宁东": ("广西", "南宁"), "南宁西": ("广西", "南宁"), "南通西": ("江苏", "南通"),
    "南阳东": ("河南", "南阳"), "原平西": ("山西", "忻州"), "双辽": ("吉林", "四平"),
    "吉安西": ("江西", "吉安"), "启东": ("江苏", "南通"), "呼和浩特": ("内蒙古", "呼和浩特"),
    "呼和浩特东": ("内蒙古", "呼和浩特"), "商丘东": ("河南", "商丘"), "嘉峪关南": ("甘肃", "嘉峪关"),
    "图们北": ("吉林", "延边"), "大同南": ("山西", "大同"), "大庆东": ("黑龙江", "大庆"),
    "大明湖": ("山东", "济南"), "天水南": ("甘肃", "天水"), "天津北": ("天津", "天津"),
    "婺源": ("江西", "上饶"), "孝感东": ("湖北", "孝感"), "安顺西": ("贵州", "安顺"),
    "宋城路": ("河南", "开封"), "宜宾西": ("四川", "宜宾"), "开封北": ("河南", "开封"),
    "开阳": ("贵州", "贵阳"), "张家界西": ("湖南", "张家界"), "怀化南": ("湖南", "怀化"),
    "成都西": ("四川", "成都"), "承德南": ("河北", "承德"), "攀枝花南": ("四川", "攀枝花"),
    "日照西": ("山东", "日照"), "昆明南": ("云南", "昆明"), "杭州南": ("浙江", "杭州"),
    "桂林北": ("广西", "桂林"), "桂林西": ("广西", "桂林"), "汉口": ("湖北", "武汉"),
    "沈阳南": ("辽宁", "沈阳"), "沙坪坝": ("重庆", "重庆"), "洛阳龙门": ("河南", "洛阳"),
    "淮南南": ("安徽", "淮南"), "湘潭北": ("湖南", "湘潭"), "潍坊北": ("山东", "潍坊"),
    "烟台南": ("山东", "烟台"), "珲春": ("吉林", "延边"), "石家庄东": ("河北", "石家庄"),
    "福田": ("广东", "深圳"), "秦皇岛": ("河北", "秦皇岛"), "荣成": ("山东", "威海"),
    "襄州": ("湖北", "襄阳"), "西安北": ("陕西", "西安"), "贵阳东": ("贵州", "贵阳"),
    "贵阳北": ("贵州", "贵阳"), "赣州西": ("江西", "赣州"), "连云港东": ("江苏", "连云港"),
    "鄂尔多斯": ("内蒙古", "鄂尔多斯"), "阜阳西": ("安徽", "阜阳"), "阿勒泰": ("新疆", "阿勒泰"),
    "阿尔山北": ("内蒙古", "兴安"), "雅安": ("四川", "雅安"), "青岛北": ("山东", "青岛"),
    "青岛西": ("山东", "青岛"), "鹰潭北": ("江西", "鹰潭"),
}

# 手工补充坐标（WGS84 近似，用于长途路过车次的方向投影）
MANUAL_COORD = {
    "南通西": (120.86, 32.01), "启东": (121.66, 31.81), "海安": (120.47, 32.57),
    "连云港东": (119.43, 34.72), "杭州南": (120.30, 30.17), "淮南南": (117.02, 32.63),
    "阜阳西": (115.79, 32.88), "义乌": (120.075, 29.307), "亳州南": (115.83, 33.75),
    "千岛湖": (118.95, 29.60), "汉口": (114.26, 30.62), "沈阳南": (123.43, 41.66),
    "海门": (121.15, 31.89), "界首南": (115.35, 33.27), "西安北": (108.94, 34.38),
    "贵阳北": (106.63, 26.74), "青岛北": (120.37, 36.15),
}

def to_min(t):
    if not t or t == "-":
        return None
    h, m = t.split(":")
    return int(h) * 60 + int(m)

def fmt(x):
    return None if x is None else x

# ---------- 1. 读取车站省份映射 ----------
prov, city = {}, {}
with open(STATION_CSV, encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        if r["站名"] not in prov:
            prov[r["站名"]] = r["省"]
            city[r["站名"]] = r["市"]
for s, (p, c) in MANUAL.items():
    prov[s] = p
    city[s] = c

def is_yrd(st):
    return prov.get(st) in YRD_PROVINCES

# ---------- 2. 读取全部停靠记录 ----------
# row: (station, type, status, train_no, src, dst, arrival, departure)
rows = []
for f in glob.glob(os.path.join(SCHED, "*.csv")):
    station = os.path.splitext(os.path.basename(f))[0]
    with open(f, encoding="gb18030", errors="replace") as fh:
        for r in csv.reader(fh):
            if len(r) >= 7:
                rows.append((station,) + tuple(r[:7]))

# ---------- 3. 按车次合并 & 去重 ----------
by_train = defaultdict(list)
for r in rows:
    by_train[r[3]].append(r)

# 同一车次同一车站出现多行（Z97/Z98/Z99/Z100 香港直通车双列）：优先保留非"往香港直通车"行
clean = {}
for t, v in by_train.items():
    seen = {}
    for r in v:
        if r[0] in seen:
            if r[1] == "往香港直通车":
                continue
            seen[r[0]] = r  # 覆盖，保留后读到的普通行
        else:
            seen[r[0]] = r
    clean[t] = list(seen.values())

# ---------- 4. 只对长三角停靠站排序（长三角内部穿越必然 <24h，时钟值在圆上单调） ----------
dep_t = lambda r: to_min(r[7]) if to_min(r[7]) is not None else to_min(r[6])
arr_t = lambda r: to_min(r[6]) if to_min(r[6]) is not None else to_min(r[7])

def load_coords():
    coords = {}
    with open(STATION_CSV, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            try:
                coords[r["站名"]] = (float(r["WGS84_Lng"]), float(r["WGS84_Lat"]))
            except (ValueError, TypeError):
                pass
    coords.update(MANUAL_COORD)
    return coords

COORDS = load_coords()

def order_yrd_stops(stops):
    """返回按运行顺序排列的长三角停靠站列表；不足2站返回('skip',)，
    无法可靠排序返回('fail',)"""
    yrd = [r for r in stops if is_yrd(r[0])]
    if len(set(r[0] for r in yrd)) < 2:
        return ("skip",)
    start = next((r for r in stops if r[2] == "始"), None)
    end = next((r for r in stops if r[2] == "终"), None)
    yrd_names = set(r[0] for r in yrd)

    # 情形A：始发站在长三角 -> 以始发时刻为锚前向排序（精确）
    if start and start[0] in yrd_names:
        t0 = dep_t(start)
        return sorted(yrd, key=lambda r: 0 if r is start else (arr_t(r) - t0) % 1440)

    # 情形B：终到站在长三角 -> 以终到时刻为锚后向排序（精确）：
    # 按"距终到的分钟数"降序 = 运行顺序
    if end and end[0] in yrd_names:
        t1 = arr_t(end)
        return sorted(yrd, key=lambda r: -((t1 - (arr_t(r) if r is end else dep_t(r))) % 1440))

    # 情形C/D：路过或无锚 -> 时钟排序的所有旋转 + 起终点方向投影打分
    clock_sorted = sorted(yrd, key=arr_t)
    n = len(clock_sorted)
    # 候选旋转：在 i 处切开（第 i 个为起点，回绕补 1440）
    def elapsed_of(order):
        vals, prev = [], None
        for r in order:
            v = arr_t(r)
            if prev is not None and v <= prev:
                v += 1440
            vals.append(v)
            prev = v
        return vals
    # 投影方向：src -> dst
    src_name, dst_name = stops[0][4], stops[0][5]
    p0, p1 = COORDS.get(src_name), COORDS.get(dst_name)
    def proj(name):
        x, y = COORDS.get(name, (None, None))
        if x is None or p0 is None or p1 is None:
            return None
        dx, dy = p1[0] - p0[0], p1[1] - p0[1]
        nrm = (dx * dx + dy * dy) ** 0.5
        return ((x - p0[0]) * dx + (y - p0[1]) * dy) / nrm if nrm > 1e-9 else None

    best, best_score = None, None
    for i in range(n):
        cand = clock_sorted[i:] + clock_sorted[:i]
        vals = elapsed_of(cand)
        span = vals[-1] - vals[0]
        gaps = [vals[k + 1] - vals[k] for k in range(n - 1)]
        # 合法性: 跨越长三角不超过 20h，相邻间隔不超过 10h，时间严格递增
        if span > 20 * 60 or any(g <= 0 or g > 10 * 60 for g in gaps):
            continue
        # 打分: 与投影方向的一致性（Spearman footrule 越小越好）
        if p0 is not None and p1 is not None:
            pr = {r[0]: proj(r[0]) for r in cand}
            if all(v is not None for v in pr.values()):
                rank_proj = {s: k for k, s in enumerate(sorted(pr, key=lambda s: pr[s]))}
                score = sum(abs(k - rank_proj[r[0]]) for k, r in enumerate(cand))
            else:
                score = span  # 无坐标兜底: 取跨度最小的旋转
        else:
            score = span
        if best_score is None or score < best_score:
            best, best_score = cand, score
    if best is None:
        return ("fail",)
    return best

train_routes = {}   # train_no -> ordered YRD stops
skipped_single, failed_order = [], []
for t, stops in clean.items():
    ordered = order_yrd_stops(stops)
    if ordered == ("skip",):
        skipped_single.append(t)
    elif ordered == ("fail",):
        failed_order.append(t)
    else:
        train_routes[t] = ordered

print(f"车次总数: {len(clean)}, 长三角内停靠>=2站: {len(train_routes)}, "
      f"仅停1站忽略: {len(skipped_single)}, 无法可靠排序忽略: {len(failed_order)}")
if failed_order:
    print("  排序失败车次:", failed_order[:20], "..." if len(failed_order) > 20 else "")

# ---------- 5. 生成 train 表（起终点裁剪规则） ----------
train_table = []      # (train_no, station_order, station_name, arrival, departure, type)
clip_stats = {"origin_clipped": 0, "dest_clipped": 0, "both_clipped": 0, "internal": 0}

for t, stops in clean.items():
    if t not in train_routes:
        continue
    dedup = train_routes[t]
    start = next((r for r in stops if r[2] == "始"), None)
    end = next((r for r in stops if r[2] == "终"), None)
    n = len(dedup)
    o_clipped = not (start and start[0] == dedup[0][0])
    d_clipped = not (end and end[0] == dedup[-1][0])
    if o_clipped and d_clipped:
        clip_stats["both_clipped"] += 1
    elif o_clipped:
        clip_stats["origin_clipped"] += 1
    elif d_clipped:
        clip_stats["dest_clipped"] += 1
    else:
        clip_stats["internal"] += 1
    for order, r in enumerate(dedup, 1):
        st, ttype = r[0], r[1]
        arrival = r[6] if r[6] != "-" else None
        departure = r[7] if r[7] != "-" else None
        if order == 1 and (o_clipped or arrival is None):
            arrival = None           # 裁剪后的起点站：到达时间置空
        if order == n and (d_clipped or departure is None):
            departure = None         # 裁剪后的终点站：出发时间置空
        train_table.append((t, order, st, arrival, departure, ttype))

kept_trains = len(train_routes)
print(f"裁剪统计: {clip_stats}")
print(f"train 表行数: {len(train_table)}")

# ---------- 5.5 时间合理性校验 ----------
anomaly = []
by_t = defaultdict(list)
for row in train_table:
    by_t[row[0]].append(row)
for t, rs in by_t.items():
    rs.sort(key=lambda x: x[1])
    prev_dep = None
    for r in rs:
        a = to_min(r[3]) if r[3] else None
        if prev_dep is not None and a is not None:
            if (a - prev_dep) % 1440 > 10 * 60:
                anomaly.append((t, r[1], r[2]))
        if r[4]:
            prev_dep = to_min(r[4])
print(f"相邻停站间隔>10h的异常记录: {len(anomaly)}")
if anomaly:
    print("  异常样例(前10):", anomaly[:10])

# ---------- 6. 车站表 ----------
used_stations = sorted(set(r[2] for r in train_table),
                       key=lambda s: (PROV_ORDER.get(prov.get(s), 9), s))
station_id = {s: 1001 + i for i, s in enumerate(used_stations)}
station_table = [(station_id[s], s, prov.get(s, "?"), city.get(s, "?")) for s in used_stations]
print(f"车站数: {len(station_table)}")

# ---------- 7. 边表（相邻停站） ----------
edges = defaultdict(set)   # (from,to) -> set(train)
for t in sorted(set(r[0] for r in train_table)):
    seq = [r[2] for r in sorted((x for x in train_table if x[0] == t), key=lambda x: x[1])]
    for a, b in zip(seq, seq[1:]):
        edges[(a, b)].add(t)
edge_table = []
for (a, b), ts in sorted(edges.items(), key=lambda kv: (station_id[kv[0][0]], station_id[kv[0][1]])):
    edge_table.append((station_id[a], a, station_id[b], b, len(ts), ",".join(sorted(ts)[:5]) + ("..." if len(ts) > 5 else "")))
print(f"边数(有向): {len(edge_table)}")

# ---------- 8. 输出 CSV ----------
with open(os.path.join(OUT, "train.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["train_no", "station_order", "station_id", "station_name", "arrival_time", "departure_time", "train_type"])
    for t, order, st, a, d, tt in train_table:
        w.writerow([t, order, station_id[st], st, a if a else "NULL", d if d else "NULL", tt])

with open(os.path.join(OUT, "station.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["station_id", "station_name", "province", "city"])
    w.writerows(station_table)

with open(os.path.join(OUT, "edge.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["from_station_id", "from_station_name", "to_station_id", "to_station_name", "train_count", "sample_trains"])
    w.writerows(edge_table)

# ---------- 9. 输出 Excel ----------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", name="微软雅黑", size=10)
BODY_FONT = Font(name="微软雅黑", size=10)
CENTER = Alignment(horizontal="center")

def add_sheet(ws, header, data):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, CENTER
    for row in data:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, wdt in enumerate([14, 14, 12, 14, 14, 14, 12]):
        ws.column_dimensions[get_column_letter(i + 1)].width = wdt
    ws.freeze_panes = "A2"

ws = wb.active
ws.title = "车次运行图"
add_sheet(ws, ["train_no", "station_order", "station_id", "station_name", "arrival_time", "departure_time", "train_type"],
          [(t, o, sid, st, a if a else "NULL", d if d else "NULL", tt)
           for t, o, st, a, d, tt in train_table
           for sid in [station_id[st]]])

ws2 = wb.create_sheet("车站")
add_sheet(ws2, ["station_id", "station_name", "province", "city"], station_table)
for i, wdt in enumerate([12, 16, 10, 12]):
    ws2.column_dimensions[get_column_letter(i + 1)].width = wdt

ws3 = wb.create_sheet("边(相邻停站)")
add_sheet(ws3, ["from_id", "from_station", "to_id", "to_station", "车次数", "示例车次"], edge_table)
for i, wdt in enumerate([10, 14, 10, 14, 10, 30]):
    ws3.column_dimensions[get_column_letter(i + 1)].width = wdt

ws4 = wb.create_sheet("说明")
notes = [
    ["长三角铁路列车运行网络 — 数据说明"],
    [],
    ["数据来源", "railwayCrawler-main/schedulers（cnrail.geogv.org 爬取的 321 个主要车站时刻表）"],
    ["长三角范围", "上海、江苏、浙江、安徽（三省一市）"],
    ["车站范围", f"共 {len(station_table)} 个位于长三角的车站（数据集仅覆盖全国主要车站，部分小站如无锡、镇江无停靠记录）"],
    ["车次范围", f"共 {kept_trains} 个车次（在长三角内停靠 >= 2 站；仅停 1 站的 {len(skipped_single)} 个车次已忽略）"],
    [],
    ["裁剪规则"],
    ["区域外始发", "以进入长三角后的第一个停靠站作为起点站（arrival_time 置 NULL）"],
    ["区域外终到", "以离开长三角前的最后一个停靠站作为终点站（departure_time 置 NULL）"],
    ["路过长三角", "同时应用以上两条规则"],
    ["仅停 1 站", "该线路忽略"],
    [],
    ["表结构"],
    ["车次运行图", "一行 = 某车次在某站的一次停靠；station_order=1 为起点站，MAX 为终点站"],
    ["车站", "station_id 为自定义唯一编号（1001 起），同城市多车站已区分（如 上海/上海南/上海虹桥）"],
    ["边(相邻停站)", "由车次运行图自动生成：同一车次相邻两个停靠站构成一条有向边"],
    [],
    ["时间说明", "时间精确到 HH:MM，假设每天运行相同；跨零点车次按运行顺序还原"],
    ["粒度说明", "停站序列仅包含数据集中的 321 个主要车站，两站之间的中间小站未记录"],
    ["异常数据", f"Z97/Z98/Z99/Z100 香港直通车双列记录已去重；无法可靠排序的车次 {len(failed_order)} 个已忽略"],
    ["生成时间", "2026-09-02"],
]
for row in notes:
    ws4.append(row)
ws4.column_dimensions["A"].width = 16
ws4.column_dimensions["B"].width = 100
ws4["A1"].font = Font(bold=True, size=13, name="微软雅黑", color="1F4E79")
for r in range(3, len(notes) + 1):
    ws4.cell(row=r, column=1).font = Font(bold=True, name="微软雅黑", size=10)
    ws4.cell(row=r, column=2).font = BODY_FONT
    ws4.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(os.path.join(OUT, "长三角铁路运行网络.xlsx"))
print("已输出:", OUT)

# ---------- 10. 抽查 ----------
print("\n=== 抽查样例 ===")
for t in ["G1", "G7", "G7541", "K8484", "Z100"]:
    if t in [r[0] for r in train_table]:
        print(f"-- {t}")
        for r in sorted((x for x in train_table if x[0] == t), key=lambda x: x[1]):
            print(f"   {r[1]:>2} {r[2]:<6} {str(r[3]):>5} {str(r[4]):>5} {r[5]}")
